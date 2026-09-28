"""
Report builder for the admin exports (Excel .xlsx and PDF).

Every report is described once in REPORTS (title, columns, filters, summary) and
rendered by two generic builders, so all exports share the same layout:

    Excel: title band, generated date + filters, styled header row, striped data
           rows, live-formula summary block, frozen header, auto-filter, print setup.
    PDF:   landscape A4, header band, styled table that repeats its header on every
           page, colour-coded statuses, summary block, page numbers.
"""
from io import BytesIO
from datetime import datetime, date, timedelta
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)

from app.models.member_model import get_all_members
from app.models.trainer_model import get_all_trainers
from app.models.class_model import get_all_classes
from app.models.payment_model import get_all_payments
from app.models.attendance_model import get_all_attendance
from app.models.booking_model import get_all_bookings

ORG_NAME = "Gym Manager"
FONT = "Arial"

NAVY = "1E293B"
SLATE = "64748B"
STRIPE = "F1F5F9"
LINE = "CBD5E1"

STATUS_COLORS = {
    "Paid": "16A34A", "Active": "16A34A", "Attended": "16A34A",
    "Pending": "B45309",
    "Booked": "1D4ED8",
    "Failed": "DC2626", "Expired": "DC2626", "Cancelled": "DC2626",
    "Inactive": "64748B",
}


# --------------------------------------------------------------------------- #
# Value helpers
# --------------------------------------------------------------------------- #
def _clean(value):
    """Turn DB values (dates, decimals, timedeltas, None) into plain values."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600:02d}:{(total % 3600) // 60:02d}"
    if isinstance(value, Decimal):
        return float(value)
    return value


def _typed(value, kind):
    value = _clean(value)
    if kind == "money":
        return float(value) if value != "" else 0.0
    if kind == "int":
        return int(value) if value != "" else ""
    return str(value)


def _g(field):
    return lambda r: r.get(field)


def _name(first, last):
    return lambda r: f"{r.get(first) or ''} {r.get(last) or ''}".strip()


def _hm(value):
    """'7:00:00' -> '07:00'"""
    if not value:
        return ""
    parts = str(value).split(":")
    if len(parts) >= 2:
        try:
            return f"{int(parts[0]):02d}:{parts[1]}"
        except ValueError:
            return str(value)
    return str(value)


def _dt(value):
    """'2026-09-24 11:52:36' -> '2026-09-24 11:52'"""
    return str(value)[:16] if value else ""


def _months(r):
    m = r.get("duration_months")
    if not m:
        return ""
    return f"{m} month" + ("" if int(m) == 1 else "s")


def _visit_duration(r):
    if not r.get("check_out_time"):
        return "In gym"
    try:
        start = datetime.strptime(str(r["check_in_time"])[:19], "%Y-%m-%d %H:%M:%S")
        end = datetime.strptime(str(r["check_out_time"])[:19], "%Y-%m-%d %H:%M:%S")
        minutes = max(int((end - start).total_seconds() // 60), 0)
        return f"{minutes // 60}h {minutes % 60:02d}m"
    except (ValueError, KeyError):
        return ""


def _class_trainer(r):
    if r.get("trainer_first_name"):
        return f"{r['trainer_first_name']} {r.get('trainer_last_name') or ''}".strip()
    return "Unassigned"


def _booking_time(r):
    if not r.get("start_time"):
        return ""
    return f"{_hm(r.get('start_time'))} - {_hm(r.get('end_time'))}"


# --------------------------------------------------------------------------- #
# Summary item constructors.
#   value(rows)    -> python value (used by PDF, and Excel when a sheet is empty)
#   formula(rng)   -> Excel formula string; rng(header) gives that column's data range
# --------------------------------------------------------------------------- #
def _count_rows(label, id_header):
    return {
        "label": label, "kind": "int",
        "value": lambda rows: len(rows),
        "formula": lambda rng: f"=COUNTA({rng(id_header)})",
    }


def _count_where(label, header, field, match):
    return {
        "label": label, "kind": "int",
        "value": lambda rows: sum(1 for r in rows if r.get(field) == match),
        "formula": lambda rng: f'=COUNTIF({rng(header)},"{match}")',
    }


def _sum_where(label, amount_header, amount_field, status_header, status_field, match):
    return {
        "label": label, "kind": "money",
        "value": lambda rows: sum(float(r.get(amount_field) or 0) for r in rows if r.get(status_field) == match),
        "formula": lambda rng: f'=SUMIFS({rng(amount_header)},{rng(status_header)},"{match}")',
    }


def _sum_all(label, header, field):
    return {
        "label": label, "kind": "int",
        "value": lambda rows: sum(int(r.get(field) or 0) for r in rows),
        "formula": lambda rng: f"=SUM({rng(header)})",
    }


def _count_in_gym(label):
    return {
        "label": label, "kind": "int",
        "value": lambda rows: sum(1 for r in rows if not r.get("check_out_time")),
        "formula": lambda rng: f'=COUNTIF({rng("Check-out")},"In gym")',
    }


# --------------------------------------------------------------------------- #
# Report definitions
# Column = (header, getter, kind, pdf_width_weight); kind: text | status | int | money
# --------------------------------------------------------------------------- #
REPORTS = {
    "members": {
        "sheet": "Members",
        "title": "Members Report",
        "fetch": get_all_members,
        "status_field": "membership_status",
        "date_field": "join_date",
        "columns": [
            ("ID", _g("member_id"), "int", 0.55),
            ("Name", _name("first_name", "last_name"), "text", 1.6),
            ("Email", _g("email"), "text", 2.3),
            ("Phone", _g("phone"), "text", 1.2),
            ("Plan", _g("membership_type"), "text", 0.9),
            ("Status", _g("membership_status"), "status", 0.9),
            ("Joined", _g("join_date"), "text", 1.0),
            ("Starts", _g("membership_start_date"), "text", 1.0),
            ("Expires", _g("membership_end_date"), "text", 1.0),
        ],
        "summary": [
            _count_rows("Total members", "ID"),
            _count_where("Active", "Status", "membership_status", "Active"),
            _count_where("Inactive", "Status", "membership_status", "Inactive"),
            _count_where("Expired", "Status", "membership_status", "Expired"),
        ],
    },
    "trainers": {
        "sheet": "Trainers",
        "title": "Trainers Report",
        "fetch": get_all_trainers,
        "columns": [
            ("ID", _g("trainer_id"), "int", 0.55),
            ("Name", _name("first_name", "last_name"), "text", 1.6),
            ("Email", _g("email"), "text", 2.3),
            ("Phone", _g("phone"), "text", 1.2),
            ("Specialization", _g("specialization"), "text", 1.5),
            ("Hired", _g("hire_date"), "text", 1.0),
        ],
        "summary": [_count_rows("Total trainers", "ID")],
    },
    "classes": {
        "sheet": "Classes",
        "title": "Classes Report",
        "fetch": get_all_classes,
        "columns": [
            ("ID", _g("class_id"), "int", 0.55),
            ("Class", _g("class_name"), "text", 1.8),
            ("Trainer", _class_trainer, "text", 1.6),
            ("Day", _g("schedule_day"), "text", 1.0),
            ("Starts", lambda r: _hm(r.get("start_time")), "text", 0.8),
            ("Ends", lambda r: _hm(r.get("end_time")), "text", 0.8),
            ("Capacity", _g("capacity"), "int", 0.8),
        ],
        "summary": [
            _count_rows("Total classes", "ID"),
            _sum_all("Total capacity (seats)", "Capacity", "capacity"),
        ],
    },
    "payments": {
        "sheet": "Payments",
        "title": "Payments Report",
        "fetch": get_all_payments,
        "status_field": "status",
        "date_field": "payment_date",
        "columns": [
            ("ID", _g("payment_id"), "int", 0.55),
            ("Member", _name("member_first_name", "member_last_name"), "text", 1.5),
            ("Amount (RWF)", _g("amount"), "money", 1.1),
            ("Method", _g("payment_method"), "text", 1.3),
            ("Reference", _g("payment_reference"), "text", 1.8),
            ("Duration", _months, "text", 0.9),
            ("Paid On", _g("payment_date"), "text", 1.0),
            ("Status", _g("status"), "status", 0.9),
            ("Valid From", _g("membership_start_date"), "text", 1.0),
            ("Valid Until", _g("membership_end_date"), "text", 1.0),
        ],
        "summary": [
            _count_rows("Total transactions", "ID"),
            _sum_where("Collected (Paid)", "Amount (RWF)", "amount", "Status", "status", "Paid"),
            _sum_where("Awaiting approval (Pending)", "Amount (RWF)", "amount", "Status", "status", "Pending"),
            _sum_where("Rejected (Failed)", "Amount (RWF)", "amount", "Status", "status", "Failed"),
        ],
    },
    "attendance": {
        "sheet": "Attendance",
        "title": "Attendance Report",
        "fetch": get_all_attendance,
        "date_field": "check_in_time",
        "columns": [
            ("ID", _g("attendance_id"), "int", 0.55),
            ("Member", _name("first_name", "last_name"), "text", 1.8),
            ("Check-in", lambda r: _dt(r.get("check_in_time")), "text", 1.4),
            ("Check-out", lambda r: _dt(r.get("check_out_time")) or "In gym", "text", 1.4),
            ("Duration", _visit_duration, "text", 1.0),
        ],
        "summary": [
            _count_rows("Total check-ins", "ID"),
            _count_in_gym("Currently in the gym"),
        ],
    },
    "bookings": {
        "sheet": "Bookings",
        "title": "Class Bookings Report",
        "fetch": get_all_bookings,
        "status_field": "status",
        "date_field": "booking_date",
        "columns": [
            ("ID", _g("booking_id"), "int", 0.55),
            ("Member", _name("member_first_name", "member_last_name"), "text", 1.6),
            ("Class", _g("class_name"), "text", 1.8),
            ("Day", _g("schedule_day"), "text", 1.0),
            ("Time", _booking_time, "text", 1.2),
            ("Booked On", _g("booking_date"), "text", 1.0),
            ("Status", _g("status"), "status", 0.9),
        ],
        "summary": [
            _count_rows("Total bookings", "ID"),
            _count_where("Booked", "Status", "status", "Booked"),
            _count_where("Attended", "Status", "status", "Attended"),
            _count_where("Cancelled", "Status", "status", "Cancelled"),
        ],
    },
}

# Order used for the full-system report
FULL_ORDER = ["members", "trainers", "classes", "payments", "attendance", "bookings"]

# Summary sheet / page of the full-system report.
#   value(data)    -> python value, data = {report_key: rows}
#   formula(ref)   -> Excel formula, ref(report_key, header) = qualified column range
FULL_SUMMARY = [
    {"label": "Total members", "kind": "int", "needs": ["members"],
     "value": lambda d: len(d["members"]),
     "formula": lambda ref: f"=COUNTA({ref('members', 'ID')})"},
    {"label": "Active memberships", "kind": "int", "needs": ["members"],
     "value": lambda d: sum(1 for r in d["members"] if r.get("membership_status") == "Active"),
     "formula": lambda ref: f'=COUNTIF({ref("members", "Status")},"Active")'},
    {"label": "Trainers", "kind": "int", "needs": ["trainers"],
     "value": lambda d: len(d["trainers"]),
     "formula": lambda ref: f"=COUNTA({ref('trainers', 'ID')})"},
    {"label": "Classes", "kind": "int", "needs": ["classes"],
     "value": lambda d: len(d["classes"]),
     "formula": lambda ref: f"=COUNTA({ref('classes', 'ID')})"},
    {"label": "Active class bookings", "kind": "int", "needs": ["bookings"],
     "value": lambda d: sum(1 for r in d["bookings"] if r.get("status") == "Booked"),
     "formula": lambda ref: f'=COUNTIF({ref("bookings", "Status")},"Booked")'},
    {"label": "Gym check-ins (all time)", "kind": "int", "needs": ["attendance"],
     "value": lambda d: len(d["attendance"]),
     "formula": lambda ref: f"=COUNTA({ref('attendance', 'ID')})"},
    {"label": "Revenue collected (RWF)", "kind": "money", "needs": ["payments"],
     "value": lambda d: sum(float(r.get("amount") or 0) for r in d["payments"] if r.get("status") == "Paid"),
     "formula": lambda ref: f'=SUMIFS({ref("payments", "Amount (RWF)")},{ref("payments", "Status")},"Paid")'},
    {"label": "Payments awaiting approval", "kind": "int", "needs": ["payments"],
     "value": lambda d: sum(1 for r in d["payments"] if r.get("status") == "Pending"),
     "formula": lambda ref: f'=COUNTIF({ref("payments", "Status")},"Pending")'},
]


# --------------------------------------------------------------------------- #
# Data collection + filtering
# --------------------------------------------------------------------------- #
def _apply_filters(report, rows, filters):
    status = filters.get("status")
    date_from = filters.get("date_from")
    date_to = filters.get("date_to")

    status_field = report.get("status_field")
    date_field = report.get("date_field")

    if status and status_field:
        rows = [r for r in rows if str(r.get(status_field)) == status]

    if date_field and (date_from or date_to):
        def day_of(r):
            return str(_clean(r.get(date_field)))[:10]
        if date_from:
            rows = [r for r in rows if day_of(r) >= date_from]
        if date_to:
            rows = [r for r in rows if day_of(r) <= date_to]
    return rows


def _filter_text(report, filters):
    parts = []
    if filters.get("status") and report.get("status_field"):
        parts.append(f"Status: {filters['status']}")
    if report.get("date_field"):
        if filters.get("date_from"):
            parts.append(f"From {filters['date_from']}")
        if filters.get("date_to"):
            parts.append(f"To {filters['date_to']}")
    return " | ".join(parts) if parts else "All records"


def collect_sections(keys, filters):
    sections = []
    for key in keys:
        report = REPORTS[key]
        rows = _apply_filters(report, report["fetch"](), filters)
        sections.append({"key": key, "report": report, "rows": rows})
    return sections


def _now_text():
    return datetime.now().strftime("%d %b %Y, %H:%M")


# --------------------------------------------------------------------------- #
# EXCEL
# --------------------------------------------------------------------------- #
def _xf(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


_THIN = Side(style="thin", color=LINE)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor=NAVY)
_STRIPE_FILL = PatternFill("solid", fgColor=STRIPE)

HEADER_ROW = 4


def _setup_print(ws, title_rows=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.left.text = ORG_NAME
    if title_rows:
        ws.print_title_rows = title_rows


def _write_table_sheet(ws, section, filters, generated):
    """Writes one report sheet. Returns info needed for cross-sheet formulas."""
    report = section["report"]
    rows = section["rows"]
    cols = report["columns"]
    n = len(cols)
    last_col = get_column_letter(n)
    letters = {c[0]: get_column_letter(i) for i, c in enumerate(cols, start=1)}

    ws.sheet_view.showGridLines = False

    # Title band
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"{ORG_NAME}  |  {report['title']}"
    ws["A1"].font = _xf(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = _HEADER_FILL
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"Generated: {generated}    |    {_filter_text(report, filters)}    |    {len(rows)} record(s)"
    ws["A2"].font = _xf(size=9, italic=True, color=SLATE)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)

    # Header row
    for c, (header, _, kind, _w) in enumerate(cols, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=header)
        cell.font = _xf(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[HEADER_ROW].height = 24

    widths = [len(c[0]) for c in cols]

    # Data rows
    for i, r in enumerate(rows):
        excel_row = HEADER_ROW + 1 + i
        for c, (header, getter, kind, _w) in enumerate(cols, start=1):
            value = _typed(getter(r), kind)
            cell = ws.cell(row=excel_row, column=c, value=value)
            cell.border = _BORDER
            cell.font = _xf()
            if i % 2 == 1:
                cell.fill = _STRIPE_FILL
            if kind == "money":
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif kind == "int":
                cell.alignment = Alignment(horizontal="center")
            elif kind == "status":
                cell.font = _xf(bold=True, color=STATUS_COLORS.get(value, "334155"))
                cell.alignment = Alignment(horizontal="center")
            widths[c - 1] = max(widths[c - 1], len(f"{value:,.0f}") if kind == "money" else len(str(value)))

    if not rows:
        ws.merge_cells(start_row=HEADER_ROW + 1, start_column=1, end_row=HEADER_ROW + 1, end_column=n)
        cell = ws.cell(row=HEADER_ROW + 1, column=1, value="No records found.")
        cell.font = _xf(italic=True, color=SLATE)
        cell.alignment = Alignment(horizontal="center")

    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 3, 9), 42)

    first = HEADER_ROW + 1
    last = HEADER_ROW + len(rows)

    if rows:
        ws.freeze_panes = ws.cell(row=first, column=1)
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last}"

    # Summary block (live formulas so totals follow any edits / deletions)
    def rng(header):
        L = letters[header]
        return f"${L}${first}:${L}${last}"

    row = HEADER_ROW + max(len(rows), 1) + 2
    ws.cell(row=row, column=1, value="Summary").font = _xf(bold=True, size=11, color=NAVY)
    row += 1
    for item in report["summary"]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        label = ws.cell(row=row, column=1, value=item["label"])
        label.font = _xf(bold=True)
        value = item["formula"](rng) if rows else item["value"](rows)
        vcell = ws.cell(row=row, column=3, value=value)
        vcell.font = _xf(bold=True)
        vcell.alignment = Alignment(horizontal="left")
        vcell.number_format = '#,##0" RWF"' if item["kind"] == "money" else "#,##0"
        row += 1

    _setup_print(ws, title_rows=f"{HEADER_ROW}:{HEADER_ROW}")

    if not rows:
        return None
    return {"sheet": report["sheet"], "letters": letters, "first": first, "last": last}


def _write_summary_sheet(ws, data, info, generated):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:B1")
    ws["A1"] = f"{ORG_NAME}  |  System Report"
    ws["A1"].font = _xf(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = _HEADER_FILL
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:B2")
    ws["A2"] = f"Generated: {generated}    |    Complete snapshot of all records"
    ws["A2"].font = _xf(size=9, italic=True, color=SLATE)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)

    for c, header in enumerate(["Metric", "Value"], start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=header)
        cell.font = _xf(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def ref(key, header):
        i = info[key]
        L = i["letters"][header]
        return f"{i['sheet']}!${L}${i['first']}:${L}${i['last']}"

    row = HEADER_ROW + 1
    for idx, item in enumerate(FULL_SUMMARY):
        label = ws.cell(row=row, column=1, value=item["label"])
        label.font = _xf()
        label.border = _BORDER
        if all(k in info for k in item["needs"]):
            value = item["formula"](ref)
        else:
            value = item["value"](data)
        vcell = ws.cell(row=row, column=2, value=value)
        vcell.font = _xf(bold=True)
        vcell.border = _BORDER
        vcell.alignment = Alignment(horizontal="right")
        vcell.number_format = '#,##0" RWF"' if item["kind"] == "money" else "#,##0"
        if idx % 2 == 1:
            label.fill = _STRIPE_FILL
            vcell.fill = _STRIPE_FILL
        row += 1

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 24
    _setup_print(ws)


def build_excel(sections, filters, full=False):
    wb = Workbook()
    wb.remove(wb.active)
    generated = _now_text()

    summary_ws = wb.create_sheet("Summary") if full else None

    info = {}
    for s in sections:
        ws = wb.create_sheet(s["report"]["sheet"])
        result = _write_table_sheet(ws, s, filters, generated)
        if result:
            info[s["key"]] = result

    if full:
        data = {s["key"]: s["rows"] for s in sections}
        _write_summary_sheet(summary_ws, data, info, generated)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf_styles():
    base = getSampleStyleSheet()
    navy = colors.HexColor("#" + NAVY)
    slate = colors.HexColor("#" + SLATE)
    return {
        "title": ParagraphStyle("rTitle", parent=base["Title"], fontName="Helvetica-Bold",
                                fontSize=20, leading=24, textColor=navy, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("rSub", parent=base["Normal"], fontName="Helvetica",
                              fontSize=9, textColor=slate, spaceAfter=10),
        "h2": ParagraphStyle("rH2", parent=base["Normal"], fontName="Helvetica-Bold",
                             fontSize=11, textColor=navy, spaceBefore=12, spaceAfter=4),
        "head": ParagraphStyle("rHead", parent=base["Normal"], fontName="Helvetica-Bold",
                               fontSize=8, leading=10, textColor=colors.white, alignment=1),
        "cell": ParagraphStyle("rCell", parent=base["Normal"], fontName="Helvetica",
                               fontSize=8, leading=10),
        "cell_right": ParagraphStyle("rCellR", parent=base["Normal"], fontName="Helvetica",
                                     fontSize=8, leading=10, alignment=2),
        "cell_center": ParagraphStyle("rCellC", parent=base["Normal"], fontName="Helvetica",
                                      fontSize=8, leading=10, alignment=1),
        "label": ParagraphStyle("rLabel", parent=base["Normal"], fontName="Helvetica-Bold",
                                fontSize=9, leading=12),
        "value": ParagraphStyle("rValue", parent=base["Normal"], fontName="Helvetica",
                                fontSize=9, leading=12),
    }


def _format_summary_value(value, kind):
    if kind == "money":
        return f"{value:,.0f} RWF"
    return f"{value:,}" if isinstance(value, int) else str(value)


def _pdf_summary_table(items, styles, width=110 * mm):
    """items: list of (label, value, kind)"""
    data = [
        [Paragraph(escape(label), styles["label"]),
         Paragraph(escape(_format_summary_value(value, kind)), styles["value"])]
        for label, value, kind in items
    ]
    table = Table(data, colWidths=[width * 0.62, width * 0.38], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + LINE)),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def _pdf_section(section, filters, styles, generated, avail_width):
    report = section["report"]
    rows = section["rows"]
    cols = report["columns"]

    story = [
        Paragraph(escape(report["title"]), styles["title"]),
        Paragraph(
            escape(f"Generated {generated}  |  {_filter_text(report, filters)}  |  {len(rows)} record(s)"),
            styles["sub"],
        ),
    ]

    if not rows:
        story.append(Paragraph("No records found.", styles["cell"]))
    else:
        total_weight = sum(c[3] for c in cols)
        col_widths = [avail_width * c[3] / total_weight for c in cols]

        data = [[Paragraph(escape(c[0]), styles["head"]) for c in cols]]
        for r in rows:
            line = []
            for header, getter, kind, _w in cols:
                value = _typed(getter(r), kind)
                if kind == "money":
                    line.append(Paragraph(escape(f"{value:,.0f}"), styles["cell_right"]))
                elif kind == "int":
                    line.append(Paragraph(escape(str(value) or "-"), styles["cell_center"]))
                elif kind == "status":
                    color = STATUS_COLORS.get(value, "334155")
                    line.append(Paragraph(
                        f'<font color="#{color}"><b>{escape(value or "-")}</b></font>',
                        styles["cell_center"]))
                else:
                    line.append(Paragraph(escape(value or "-"), styles["cell"]))
            data.append(line)

        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + STRIPE)]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + LINE)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)

    items = [(i["label"], i["value"](rows), i["kind"]) for i in report["summary"]]
    story.append(KeepTogether([
        Paragraph("Summary", styles["h2"]),
        _pdf_summary_table(items, styles),
    ]))
    return story


def _make_page_decorator(generated):
    page_w, page_h = landscape(A4)

    def decorate(canvas, doc):
        canvas.saveState()
        # top band
        canvas.setFillColor(colors.HexColor("#" + NAVY))
        canvas.rect(0, page_h - 9 * mm, page_w, 9 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(15 * mm, page_h - 6 * mm, ORG_NAME.upper())
        # footer
        canvas.setFillColor(colors.HexColor("#94A3B8"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(15 * mm, 8 * mm, f"{ORG_NAME}  |  Confidential  |  Generated {generated}")
        canvas.drawRightString(page_w - 15 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    return decorate


def build_pdf(sections, filters, full=False):
    buf = BytesIO()
    generated = _now_text()
    page_w, _ = landscape(A4)
    margin = 15 * mm

    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=margin, rightMargin=margin, topMargin=18 * mm, bottomMargin=16 * mm,
        title=f"{ORG_NAME} Report", author=ORG_NAME,
    )
    styles = _pdf_styles()
    avail = page_w - 2 * margin

    story = []
    if full:
        data = {s["key"]: s["rows"] for s in sections}
        items = [(i["label"], i["value"](data), i["kind"]) for i in FULL_SUMMARY]
        story += [
            Paragraph("System Report", styles["title"]),
            Paragraph(escape(f"Generated {generated}  |  Complete snapshot of all records"), styles["sub"]),
            Paragraph("Overview", styles["h2"]),
            _pdf_summary_table(items, styles),
        ]

    for i, s in enumerate(sections):
        if full or i > 0:
            story.append(PageBreak())
        story += _pdf_section(s, {} if full else filters, styles, generated, avail)

    decorate = _make_page_decorator(generated)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    buf.seek(0)
    return buf